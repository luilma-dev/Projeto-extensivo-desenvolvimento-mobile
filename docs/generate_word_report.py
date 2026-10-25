import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def style_table(table, data, DARK):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        is_header = (i == 0)
        for j, cell in enumerate(row.cells):
            cell.text = data[i][j]
            set_cell_margins(cell, top=75, bottom=75, left=95, right=95)
            if is_header:
                set_cell_background(cell, "0066CC")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True
                        r.font.size = Pt(9)
                        r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "F8FAFC" if i % 2 == 1 else "FFFFFF")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8.5)
                        r.font.color.rgb = DARK

def create_report():
    doc = Document()
    
    # Page setup (margins: ~2.3cm)
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
    
    # Base colors
    PRIMARY = RGBColor(0, 102, 204)    # Blue
    SECONDARY = RGBColor(70, 80, 95)   # Muted slate
    DARK = RGBColor(15, 23, 42)        # Ink
    
    # Header Info
    p_meta = doc.add_paragraph()
    r_meta1 = p_meta.add_run("FÁBRICA DE SOFTWARE — PROGRAMAÇÃO DE DISPOSITIVOS MÓVEIS EM ANDROID\n")
    r_meta1.bold = True
    r_meta1.font.size = Pt(10)
    r_meta1.font.color.rgb = SECONDARY
    r_meta2 = p_meta.add_run("Professor: Fábio Ferreira  |  Outubro / 2026\n")
    r_meta2.font.size = Pt(9.5)
    r_meta2.font.color.rgb = SECONDARY
    p_meta.paragraph_format.space_after = Pt(12)
    
    # Main Title
    p_title = doc.add_paragraph()
    r_title = p_title.add_run("Relatório de Extensão — Item 3: Planejamento e Desenvolvimento do Projeto")
    r_title.bold = True
    r_title.font.size = Pt(20)
    r_title.font.color.rgb = PRIMARY
    p_title.paragraph_format.space_after = Pt(4)
    
    # Subtitle
    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("Atividade Prática 3: Sprint 3 — Persistência Local (AsyncStorage e SQLite), Consumo de APIs REST (HTTPS) e Arquitetura Offline-First\nProjeto: TruckFlow — Sistema Móvel para Oficinas de Caminhões")
    r_sub.font.size = Pt(11)
    r_sub.italic = True
    r_sub.font.color.rgb = SECONDARY
    p_sub.paragraph_format.space_after = Pt(16)
    
    # Team Callout Box
    table_team = doc.add_table(rows=1, cols=1)
    table_team.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_team = table_team.rows[0].cells[0]
    cell_team.width = Inches(6.7)
    set_cell_background(cell_team, "F1F5F9")
    set_cell_margins(cell_team, top=140, bottom=140, left=180, right=180)
    p_box = cell_team.paragraphs[0]
    r_box_title = p_box.add_run("Equipe Responsável (Sprint 3):\n")
    r_box_title.bold = True
    r_box_title.font.size = Pt(10)
    r_box_title.font.color.rgb = PRIMARY
    p_box.add_run("• Cauê — Persistência Local Chave-Valor (AsyncStorage), Sessão e Autenticação (preferences.ts, LoginScreen.tsx)\n")
    p_box.add_run("• Cauã — Banco de Dados Relacional Embarcado (SQLite), DDL, CRUD e Transações (database.ts, database.web.ts)\n")
    p_box.add_run("• Alexandre — Consumo de APIs REST via HTTPS (Tabela FIPE), Cliente HTTP e Feedback Visual (client.ts, fipe.ts, FipePicker.tsx)\n")
    p_box.add_run("• Vinícius — Arquitetura Offline-First, Monitoramento de Rede (NetInfo) e Sincronização Outbox (sync.ts, App.tsx, HomeScreen.tsx)\n")
    r_hint = p_box.add_run("(Turma: Fábrica de Software — Programação de Dispositivos Móveis em Android)")
    r_hint.font.size = Pt(8.5)
    r_hint.italic = True
    r_hint.font.color.rgb = SECONDARY
    
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = PRIMARY
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(11.5)
        run.font.color.rgb = DARK
        return p

    def add_body(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.size = Pt(10)
        run.font.color.rgb = DARK
        return p

    def add_caption(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(10)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.size = Pt(9)
        run.italic = True
        run.font.color.rgb = SECONDARY
        return p

    # --- 3.1 VISÃO GERAL ---
    add_heading_1("3.1 Visão Geral da Sprint 3")
    add_body("Na Sprint 3 (laboratórios de Outubro/2026), o aplicativo móvel TruckFlow passou por uma evolução arquitetural fundamental: a transição de um protótipo puramente visual (cujos dados residiam apenas em memória volátil via useState) para uma aplicação móvel inteligente, resiliente e integrada a dados físicos e remotos via HTTPS.")
    add_body("O público-alvo do aplicativo — mecânicos, gestores e colaboradores de oficinas de veículos pesados e transportadoras — opera frequentemente em galpões industriais, rodovias ou pátios de carga com conectividade de rede celular instável ou inexistente. Por esta razão, os 4 integrantes implementaram os quatro pilares dos Temas 4 e 5 do SAVA de forma coordenada:")
    add_body("1. Persistência Local Chave-Valor com AsyncStorage (Cauê): para dados leves de configuração, retenção de sessão administrativa, histórico da última aba ativa e cache de marcas;\n"
             "2. Banco de Dados Relacional Embarcado com SQLite (Cauã): para persistência não-volátil do modelo de negócios (tabelas de veículos, ordens de serviço e fila de envio com transações ACID);\n"
             "3. Consumo de APIs REST via HTTPS com Fetch API e AbortController (Alexandre): integração segura via HTTPS com a API pública da Tabela FIPE (leitura GET), controle de timeout de 10s, tratamento de erros HTTP e indicador visual de carregamento;\n"
             "4. Arquitetura Offline-First e Sincronização Outbox (Vinícius): capacidade operacional plena sem sinal de internet, detecção automática de conexão via NetInfo e envio sequencial POST HTTPS para o servidor remoto.")

    # --- 3.2 STACK TECNOLÓGICA ---
    add_heading_1("3.2 Stack Tecnológica Adotada")
    add_body("A tabela a seguir consolida todas as tecnologias e endpoints HTTPS empregados na Sprint 3 e suas respectivas justificativas de uso no aplicativo TruckFlow:")
    
    stack_data = [
        ("Camada", "Tecnologia", "Versão / Pacote / URL HTTPS", "Finalidade no TruckFlow"),
        ("Linguagem", "TypeScript", "~6.0.3", "Tipagem estática estrita dos modelos de domínio (Vehicle, Order, FipeItem, SyncItem)."),
        ("Core Mobile", "React Native / Expo", "RN 0.86.3 / SDK 57", "Desenvolvimento cross-platform nativo para Android com suporte a Expo Go."),
        ("Persistência Chave-Valor", "AsyncStorage", "@react-native-async-storage/async-storage", "Armazenamento leve: sessão, última aba ativa, timestamp de sincronização e cache offline."),
        ("Banco Embarcado", "SQLite Nativo", "expo-sqlite ~57.0.4", "Armazenamento relacional físico (ACID) com chaves estrangeiras, WAL e transações."),
        ("Cliente HTTP/HTTPS", "Fetch API + AbortController", "ES2024 Nativo (TLS/HTTPS)", "Requisições HTTPS com cabeçalhos JSON, timeout de 10s e classe customizada ApiError."),
        ("API Pública HTTPS (GET)", "Tabela FIPE REST", "https://parallelum.com.br/fipe/api/v1/caminhoes", "Consulta em tempo real de montadoras e modelos oficiais de caminhões pesados."),
        ("API Remota HTTPS (POST)", "JSONPlaceholder REST", "https://jsonplaceholder.typicode.com/posts", "Endpoint REST de alta disponibilidade para validação da fila outbox (POST -> 201 Created)."),
        ("Gestão de Rede", "NetInfo", "@react-native-community/netinfo 12.0.1", "Monitoramento reativo da conectividade e disparo automático da sincronização."),
    ]
    
    table_stack = doc.add_table(rows=len(stack_data), cols=4)
    style_table(table_stack, stack_data, DARK)

    add_body("")
    add_heading_2("Justificativa da Coexistência: AsyncStorage e SQLite")
    add_body("A arquitetura do TruckFlow aplica o princípio da responsabilidade única:\n"
             "• AsyncStorage (implementado por Cauê) é utilizado para pequenos pares chave-valor sem integridade relacional (sessão autenticada, última aba aberta pelo usuário, data/hora da última sincronização bem-sucedida e cache offline da lista de marcas da Tabela FIPE);\n"
             "• SQLite (implementado por Cauã) é utilizado para a estrutura crítica do negócio: integridade referencial (uma ordem de serviço depende da chave primária do veículo através da FK vehicle_id), transações atômicas com rollback em caso de falha e fila persistente outbox.")

    # --- 3.3 ESTRUTURA MODULAR DO PROJETO ---
    add_heading_1("3.3 Estrutura Modular do Projeto")
    add_body("O código-fonte da Sprint 3 foi reorganizado em módulos desacoplados para atender integralmente aos requisitos de 'módulo de persistência local' e 'módulo de conexão remota HTTPS':")
    
    tree_text = (
        "TruckFlow/\n"
        "├── App.tsx                     # Orquestrador central: inicialização, NetInfo e ciclo de sync (Vinícius)\n"
        "├── src/\n"
        "│   ├── storage/                # MÓDULO DE PERSISTÊNCIA LOCAL (Cauê e Cauã)\n"
        "│   │   ├── preferences.ts      # AsyncStorage: sessão, última aba, sync timestamp, cache FIPE (Cauê)\n"
        "│   │   ├── database.ts         # SQLite nativo: abertura, DDL, transações, CRUD, outbox (Cauã)\n"
        "│   │   └── database.web.ts     # Fallback web homologado para testes em navegadores (Cauã)\n"
        "│   ├── api/                    # MÓDULO DE CONEXÃO REMOTA HTTPS (Alexandre e Vinícius)\n"
        "│   │   ├── client.ts           # Cliente HTTPS base: Fetch + timeout 10s + AbortController (Alexandre)\n"
        "│   │   ├── fipe.ts             # Consumo GET HTTPS: Marcas e modelos na Tabela FIPE (Alexandre)\n"
        "│   │   └── sync.ts             # Consumo POST HTTPS: Despacho sequencial da fila outbox (Vinícius)\n"
        "│   ├── screens/                # CAMADA DE APRESENTAÇÃO\n"
        "│   │   ├── LoginScreen.tsx     # Login administrativo integrado à sessão do AsyncStorage (Cauê)\n"
        "│   │   ├── FleetScreen.tsx     # Cadastro/busca da frota no SQLite + seletor FIPE (Cauã / Alexandre)\n"
        "│   │   ├── OrdersScreen.tsx    # Abertura e avanço de OS com cálculo em centavos (Cauã / Vinícius)\n"
        "│   │   └── HomeScreen.tsx      # Painel financeiro e card de status de sincronização (Vinícius)\n"
        "│   ├── FipePicker.tsx          # Seletor FIPE com ActivityIndicator para consulta HTTPS (Alexandre)\n"
        "│   ├── components.tsx          # Componentes visuais reutilizáveis (botões, campos, badges)\n"
        "│   ├── data.ts                 # Tipos TypeScript, enums de status e dados seed\n"
        "│   └── styles.ts               # Design system e tokens visuais centralizados\n"
        "└── evidencias/                 # Evidências em imagem da persistência e da integração HTTPS REST"
    )
    p_tree = doc.add_paragraph()
    r_tree = p_tree.add_run(tree_text)
    r_tree.font.name = "Consolas"
    r_tree.font.size = Pt(8.5)
    r_tree.font.color.rgb = DARK
    p_tree.paragraph_format.space_before = Pt(4)
    p_tree.paragraph_format.space_after = Pt(12)

    # --- 3.4 ARQUITETURA OFFLINE-FIRST, PERSISTÊNCIA E COMUNICAÇÃO HTTPS ---
    add_heading_1("3.4 Arquitetura de Dados Offline-First, Persistência e Comunicação HTTPS")
    add_body("Na arquitetura Offline-First do TruckFlow, o banco de dados local embarcado no celular é a fonte única da verdade (Single Source of Truth). Toda operação de escrita é gravada de imediato no SQLite em uma transação ACID, registrando simultaneamente um evento na tabela sync_queue (padrão Outbox). A tela atualiza instantaneamente e o módulo de rede envia os dados via HTTPS quando houver conectividade.")

    evid_dir = r"C:\Truckflow\Projeto-extensivo-desenvolvimento-mobile-main\evidencias"
    
    # Diagram 1
    add_heading_2("Diagrama 1: Arquitetura Geral Offline-First")
    d1_path = os.path.join(evid_dir, "diagrama_01_arquitetura_offline_first.png")
    if os.path.exists(d1_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        doc.add_picture(d1_path, width=Inches(6.2))
        add_caption("Figura 1 — Arquitetura de dados Offline-First: interação entre SQLite, AsyncStorage, NetInfo e APIs HTTPS REST.")

    # Diagram 2
    add_heading_2("Diagrama 2: Fluxo Sequencial de Sincronização e Resiliência")
    d2_path = os.path.join(evid_dir, "diagrama_02_fluxo_sincronizacao.png")
    if os.path.exists(d2_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        doc.add_picture(d2_path, width=Inches(6.2))
        add_caption("Figura 2 — Diagrama de Sequência: gravação atômica local e sincronização HTTPS condicional ao estado de rede.")

    # Diagram 3
    add_heading_2("Diagrama 3: Modelo Entidade-Relacionamento do Banco SQLite")
    d3_path = os.path.join(evid_dir, "diagrama_03_modelo_dados_er.png")
    if os.path.exists(d3_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        doc.add_picture(d3_path, width=Inches(5.0))
        add_caption("Figura 3 — Esquema Relacional do SQLite: tabelas de domínio (vehicles, orders) e tabela outbox (sync_queue).")

    # --- SUBSEÇÃO PADRONIZADA 1: CAUÊ ---
    add_heading_2("3.4.1 Módulo de Persistência Chave-Valor — AsyncStorage (Responsável: Cauê)")
    add_body("Implementado em src/storage/preferences.ts e integrado ao LoginScreen.tsx e ao App.tsx, este módulo encapsula a leitura e escrita assíncrona de objetos JSON no AsyncStorage com tratamento de exceções (try/catch), garantindo que preferências do operador e estado de sessão sobrevivam ao fechamento completo do aplicativo.")
    keys_data = [
        ("Chave (Key)", "Tipo de Dado (JSON)", "Descrição Técnica e Finalidade no App"),
        ("@truckflow:session", "JSON { email, loggedAt }", "Guarda as credenciais ativas do operador, mantendo a autenticação ao fechar e reabrir o app."),
        ("@truckflow:lastTab", "String ('Início'|'Frota'|'Ordens')", "Preserva o contexto de navegação: o app reabre exatamente na aba onde o usuário estava."),
        ("@truckflow:lastSync", "String ISO8601", "Timestamp da última sincronização HTTPS concluída com sucesso com o servidor remoto."),
        ("@truckflow:fipeBrands", "JSON Array<FipeItem>", "Cache não-volátil da lista de montadoras da Tabela FIPE para viabilizar consultas offline."),
    ]
    table_keys = doc.add_table(rows=len(keys_data), cols=3)
    style_table(table_keys, keys_data, DARK)

    # --- SUBSEÇÃO PADRONIZADA 2: CAUÃ ---
    add_heading_2("3.4.2 Módulo de Banco de Dados Relacional Embarcado — SQLite (Responsável: Cauã)")
    add_body("Implementado em src/storage/database.ts (com fallback homologado para web em src/storage/database.web.ts), gerencia o banco físico truckflow.db utilizando PRAGMA journal_mode = WAL e PRAGMA foreign_keys = ON. Na primeira execução, popula automaticamente os registros iniciais (seed). Todas as operações de inserção, atualização de status e exclusão executam dentro de transações atômicas (db.withTransactionAsync), gravando o dado principal e a respectiva entrada na tabela sync_queue na mesma operação indivisível.")

    # --- SUBSEÇÃO PADRONIZADA 3: ALEXANDRE ---
    add_heading_2("3.4.3 Módulo de Consumo de APIs REST via HTTPS e Feedback Visual (Responsável: Alexandre)")
    add_body("Implementado em src/api/client.ts, src/api/fipe.ts e src/FipePicker.tsx, este módulo concentra a comunicação externa sob protocolo seguro HTTPS (TLS):\n"
             "• Cliente HTTPS Base (src/api/client.ts): encapsula a Fetch API nativa definindo os cabeçalhos HTTP 'Accept: application/json' e 'Content-Type: application/json'. Utiliza a interface nativa AbortController para impor um timeout rigoroso de 10.000 ms (10 segundos), abortando requisições pendentes em redes móveis instáveis e lançando instâncias tipadas da classe ApiError.\n"
             "• Serviço REST da Tabela FIPE (src/api/fipe.ts): consome via GET os endpoints HTTPS públicos da Tabela FIPE para caminhões pesados. Ao buscar as montadoras com sucesso, atualiza o cache local no AsyncStorage; caso a requisição HTTPS falhe por ausência de rede, recupera automaticamente a lista do cache local.\n"
             "• Componente FipePicker com Feedback Visual (src/FipePicker.tsx): acoplado ao modal de cadastro em FleetScreen.tsx, renderiza o indicador nativo ActivityIndicator (spinner giratório) com o rótulo 'Consultando a Tabela FIPE...' durante toda a chamada assíncrona HTTPS.")
    
    https_data = [
        ("Método", "Endpoint HTTPS Completo", "Cabeçalhos HTTP", "Status / Resposta JSON", "Finalidade no Aplicativo"),
        ("GET", "https://parallelum.com.br/fipe/api/v1/caminhoes/marcas", "Accept: application/json", "200 OK — [{ codigo, nome }]", "Busca todas as montadoras de caminhões (Scania, Volvo, Mercedes-Benz, etc.) e atualiza o cache local."),
        ("GET", "https://parallelum.com.br/fipe/api/v1/caminhoes/marcas/{id}/modelos", "Accept: application/json", "200 OK — { modelos: [{ codigo, nome }] }", "Busca os modelos oficiais da montadora selecionada e preenche o formulário de veículo."),
        ("POST", "https://jsonplaceholder.typicode.com/posts", "Content-Type: application/json", "201 Created — { id, entity, action, data }", "Recebe via HTTPS os lotes JSON de mutações pendentes na fila outbox (sync_queue)."),
    ]
    table_https = doc.add_table(rows=len(https_data), cols=5)
    style_table(table_https, https_data, DARK)

    # --- SUBSEÇÃO PADRONIZADA 4: VINÍCIUS ---
    add_heading_2("3.4.4 Módulo de Arquitetura Offline-First, NetInfo e Sincronização Outbox (Responsável: Vinícius)")
    add_body("Implementado em src/api/sync.ts, App.tsx e src/screens/HomeScreen.tsx, garante a resiliência operacional da oficina sem sinal de internet:\n"
             "• Monitoramento Reativo de Rede (NetInfo): um listener contínuo no App.tsx monitora isConnected e isInternetReachable. Quando o aparelho fica offline, exibe o banner amarelo de alerta com o número de alterações pendentes na fila local.\n"
             "• Despacho Sequencial Outbox (syncPending em src/api/sync.ts): quando a conexão é restabelecida (ou quando o usuário aciona 'Sincronizar agora' na tela Início), percorre os itens da sync_queue enviando cada um via POST HTTPS para https://jsonplaceholder.typicode.com/posts. O registro só é removido do SQLite após confirmação 201 Created, assegurando perda zero de dados.")

    # --- 3.5 EVIDÊNCIAS DO PROTÓTIPO ---
    add_heading_1("3.5 Evidências Visuais e Comprovante de Persistência")
    add_body("Conforme solicitado nas diretrizes da Atividade Prática 3, a seguir são apresentadas as evidências visuais registradas na execução da aplicação:")
    
    evidence_items = [
        ("evidencia_01_consulta_api_fipe.png", "Evidência 1: Consumo de API REST via HTTPS (Tabela FIPE) no Cadastro de Veículos",
         "Demonstração do formulário de novo veículo integrando a API REST HTTPS pública da Tabela FIPE (https://parallelum.com.br/fipe/api/v1/caminhoes). O usuário aciona o botão 'Buscar modelo na Tabela FIPE', a requisição GET é executada exibindo indicador de atividade (ActivityIndicator) e os dados oficiais (SCANIA R-440 A 4x2) preenchem automaticamente o formulário."),
        
        ("evidencia_02_veiculo_cadastrado_persistente.png", "Evidência 2: Gravação e Leitura no Banco Relacional (SQLite)",
         "Veículo cadastrado (placa ABC9E99, Scania R-440) inserido fisicamente na tabela 'vehicles' do banco de dados relacional embarcado e exibido em tempo real na lista de veículos da oficina."),
        
        ("evidencia_03_resumo_sincronizacao_online.png", "Evidência 3: Painel de Sincronização HTTPS e Arquitetura de Rede",
         "Card central de 'Sincronização' na tela Início: exibe o badge visual '● Online' em verde, confirma que todos os dados locais foram despachados via POST HTTPS ao servidor remoto (0 pendências) e registra a data/hora da última sincronização bem-sucedida gravada no AsyncStorage."),
         
        ("evidencia_04_persistencia_apos_reload.png", "Evidência 4: Comprovante de Persistência Após Reinício do App",
         "Comprovação de não-volatilidade: o aplicativo foi completamente fechado e recarregado. Ao reabrir, a sessão administrativa permaneceu conectada, a aba ativa foi preservada e os veículos cadastrados continuam integralmente visíveis e persistidos."),
         
        ("evidencia_05_login_autenticado.png", "Evidência 5: Tela de Login Administrativo e Sessão Local",
         "Interface de autenticação com persistência de sessão: o operador realiza o login com as credenciais administrativas (admin@gmail.com / 123456) e suas preferências de acesso passam a ser gerenciadas de forma persistente."),
    ]
    
    for filename, title, desc in evidence_items:
        img_path = os.path.join(evid_dir, filename)
        add_heading_2(title)
        add_body(desc)
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(4)
            doc.add_picture(img_path, width=Inches(3.2))
            add_caption(f"Registro Visual: {title}")

    # --- 3.6 DIVISÃO DE TAREFAS PADRONIZADA ---
    add_heading_1("3.6 Divisão de Tarefas e Entregas Técnicas por Integrante (Sprint 3)")
    add_body("A equipe dividiu o desenvolvimento da Sprint 3 em 4 frentes técnicas padronizadas, garantindo que cada integrante implementasse um pilar completo da arquitetura:")
    
    tasks_data = [
        ("Integrante", "Módulo de Responsabilidade", "Arquivos Principais", "Detalhamento Técnico do Que Foi Implementado"),
        ("Cauê", "Persistência Chave-Valor (AsyncStorage) & Autenticação", "src/storage/preferences.ts, src/screens/LoginScreen.tsx", "Implementação dos métodos assíncronos de serialização JSON no AsyncStorage; retenção da sessão de login (@truckflow:session), restauração da última aba ativa (@truckflow:lastTab), timestamp de sync (@truckflow:lastSync) e cache de montadoras (@truckflow:fipeBrands)."),
        ("Cauã", "Banco de Dados Relacional Embarcado (SQLite)", "src/storage/database.ts, src/storage/database.web.ts, src/data.ts", "Modelagem DDL das tabelas relacionais (vehicles, orders, sync_queue) com PRAGMA foreign_keys = ON e modo WAL; implementação das queries SQL parametrizadas, exclusão em cascata e transações atômicas (withTransactionAsync)."),
        ("Alexandre", "Consumo de APIs REST via HTTPS (Tabela FIPE) & Feedback Visual", "src/api/client.ts, src/api/fipe.ts, src/FipePicker.tsx, src/screens/FleetScreen.tsx", "Desenvolvimento do cliente HTTP/HTTPS com Fetch API, cabeçalhos JSON, controle de timeout de 10s via AbortController e classe ApiError; consumo GET dos endpoints HTTPS da Tabela FIPE de caminhões; componente FipePicker com ActivityIndicator e fallback para cache offline."),
        ("Vinícius", "Arquitetura Offline-First, NetInfo & Sincronização Outbox", "src/api/sync.ts, App.tsx, src/screens/HomeScreen.tsx, src/screens/OrdersScreen.tsx", "Implementação do motor de sincronização outbox (syncPending) via POST HTTPS para https://jsonplaceholder.typicode.com/posts; monitoramento reativo de conectividade com NetInfo; banner de modo offline, card de sincronização na tela Início e documentação técnica."),
    ]
    table_tasks = doc.add_table(rows=len(tasks_data), cols=4)
    style_table(table_tasks, tasks_data, DARK)

    # --- 3.7 PRÓXIMOS PASSOS ---
    add_heading_1("3.7 Próximos Passos (Planejamento de Novembro)")
    add_body("Com a conclusão da Sprint 3 no dia 27/10/2026, todo o ciclo fundamental de engenharia de dados (persistência local + sincronização remota via HTTPS) está consolidado e funcional. O planejamento da equipe para o mês de Novembro contempla:")
    add_body("1. Refinamento de UI/UX: ajustes finos de ergonomia visual, tipografia e feedback tátil em testes com mecânicos reais;\n"
             "2. Simulado 2 de Apresentação: alinhamento dos discursos individuais dos 4 integrantes conforme o roteiro acadêmico;\n"
             "3. Preparação para o Pitch com a Banca Externa (24/11/2026): consolidação dos diferenciais de mercado do TruckFlow diante de bancas acadêmicas e empresariais.")

    # --- 3.8 RESUMO DIRETO AO PONTO PADRONIZADO POR INTEGRANTE ---
    add_heading_1("3.8 Resumo Direto ao Ponto: O Que Cada Integrante Implementou e Para Que Serve")
    
    add_heading_2("1. Cauê — AsyncStorage (Persistência Chave-Valor e Sessão)")
    add_body("• Arquivos: src/storage/preferences.ts e src/screens/LoginScreen.tsx\n"
             "• O que é: Armazenamento interno simples do celular para salvar textos e configurações leves no formato chave-valor (JSON).\n"
             "• Para que serve no app:\n"
             "  - Sessão (@truckflow:session): Lembra o login para não pedir e-mail e senha a cada inicialização.\n"
             "  - Navegação (@truckflow:lastTab): Lembra a última aba ativa (Início, Frota ou Ordens).\n"
             "  - Sincronização (@truckflow:lastSync): Guarda a data/hora exata do último envio remoto.\n"
             "  - Cache Offline (@truckflow:fipeBrands): Mantém a lista de montadoras da FIPE acessível mesmo sem internet.")

    add_heading_2("2. Cauã — SQLite / expo-sqlite (Banco de Dados Relacional Embarcado)")
    add_body("• Arquivos: src/storage/database.ts, src/storage/database.web.ts e src/data.ts\n"
             "• O que é: Banco de dados relacional embarcado no próprio celular (truckflow.db) com suporte a transações ACID.\n"
             "• Para que serve no app:\n"
             "  - Veículos (vehicles): Grava placa, modelo e cliente na memória física do aparelho.\n"
             "  - Ordens de Serviço (orders): Grava serviço, valores em centavos (total_cents, cost_cents), lucro e status.\n"
             "  - Integridade Referencial: Amarra ordens ao veículo (vehicle_id), excluindo em cascata se o veículo for removido.\n"
             "  - Transações Atômicas (withTransactionAsync): Insere o dado e registra na fila outbox na mesma transação segura.")

    add_heading_2("3. Alexandre — Consumo de APIs REST via HTTPS e Tabela FIPE")
    add_body("• Arquivos: src/api/client.ts, src/api/fipe.ts, src/FipePicker.tsx e src/screens/FleetScreen.tsx\n"
             "• O que é: Módulo de conexão HTTP/HTTPS (Fetch API + AbortController) para consultar servidores externos na internet.\n"
             "• Para que serve no app:\n"
             "  - Consulta HTTPS de Montadoras (GET): Busca as marcas oficiais em https://parallelum.com.br/fipe/api/v1/caminhoes/marcas.\n"
             "  - Consulta HTTPS de Modelos (GET): Busca os modelos da marca escolhida em https://parallelum.com.br/fipe/api/v1/caminhoes/marcas/{codigo}/modelos.\n"
             "  - Indicador de Carregamento (ActivityIndicator): Exibe um círculo giratório ('Consultando a Tabela FIPE...') enquanto a requisição HTTPS ocorre.\n"
             "  - Timeout de 10s e ApiError: Cancela conexões lentas após 10 segundos via AbortController para o app não travar e usa o cache local caso esteja sem internet.")

    add_heading_2("4. Vinícius — Arquitetura Offline-First, NetInfo e Sincronização Outbox")
    add_body("• Arquivos: src/api/sync.ts, App.tsx, src/screens/HomeScreen.tsx e src/screens/OrdersScreen.tsx\n"
             "• O que é: Sistema projetado para o aplicativo funcionar 100% sem sinal de internet e sincronizar quando a rede voltar.\n"
             "• Para que serve no app:\n"
             "  - Gravação Local Imediata: Toda escrita salva primeiro no SQLite local, sem esperar a internet.\n"
             "  - Fila de Envio (sync_queue): Guarda todas as alterações pendentes de envio.\n"
             "  - Monitoramento de Rede (NetInfo): Detecta queda de conexão e mostra o banner 'Modo offline'; quando a internet volta, dispara o envio automático.\n"
             "  - Sincronização Remota (POST HTTPS): Envia os itens da fila para https://jsonplaceholder.typicode.com/posts e só apaga da fila após receber 201 Created.")

    add_heading_2("5. Os 3 Diagramas Técnicos")
    add_body("• Diagrama 1 (Arquitetura Geral): Mostra a divisão entre telas, SQLite local, AsyncStorage, NetInfo e as APIs HTTPS.\n"
             "• Diagrama 2 (Fluxo Sequencial): Mostra a ordem exata das ações ao salvar (salva no banco local -> atualiza a tela -> avalia conexão -> despacha via POST HTTPS).\n"
             "• Diagrama 3 (Modelo ER): Mostra as tabelas do banco SQLite (vehicles, orders, sync_queue) e seus vínculos de chaves estrangeiras.")

    # Save to all target paths
    targets = [
        r"C:\Truckflow\Projeto-extensivo-desenvolvimento-mobile-main\docs\Modelo_do_Relatrio_de_Extensao_Sprint3.docx",
        r"C:\Truckflow\docs\Modelo_do_Relatrio_de_Extensao_Sprint3.docx",
        r"C:\Truckflow\Modelo_do_Relatrio_de_Extensao.docx",
        r"C:\Truckflow\Projeto-extensivo-desenvolvimento-mobile-main\Modelo_do_Relatrio_de_Extensao.docx",
    ]
    for target in targets:
        os.makedirs(os.path.dirname(target), exist_ok=True)
        doc.save(target)
        print(f"Relatorio gerado em: {target} ({os.path.getsize(target)} bytes)")

if __name__ == "__main__":
    create_report()
